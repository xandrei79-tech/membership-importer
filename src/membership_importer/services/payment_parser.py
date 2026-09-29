"""Parse bank statement data into normalized payment objects."""

from __future__ import annotations

import csv
from collections.abc import Iterable, Mapping, Sequence
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from ..models.payment import Payment


class PaymentParser:
    """Convert raw bank-statement rows into normalized Payment models."""

    def parse_statements(
        self,
        statement_paths: Sequence[Path | str],
    ) -> list[Payment]:
        """Read each selected bank statement and return all normalized payments."""
        payments: list[Payment] = []
        for statement_path in statement_paths:
            payments.extend(self.parse_statement(Path(statement_path)))
        return payments

    def parse_statement(
        self,
        statement: Path | str | Mapping[str, Any] | Iterable[Mapping[str, Any]],
    ) -> list[Payment]:
        """Parse a file path or a raw record collection into Payment objects."""
        if isinstance(statement, (str, Path)):
            return self._parse_file(Path(statement))
        if isinstance(statement, Mapping):
            return [self._build_payment(statement, "Unknown bank")]
        if isinstance(statement, Iterable):
            return [self._build_payment(record, "Unknown bank") for record in statement]
        raise TypeError("Unsupported statement format.")

    def _parse_file(self, path: Path) -> list[Payment]:
        """Read a statement file into a list of normalized payment objects."""
        if not path.exists():
            raise FileNotFoundError(f"Bank statement not found: {path}")

        suffix = path.suffix.lower()
        if suffix == ".csv":
            return self._parse_csv(path)
        if suffix in {".xlsx", ".xlsm", ".xls"}:
            return self._parse_excel(path)
        raise ValueError(f"Unsupported bank statement format: {path.suffix or 'unknown'}")

    def _parse_csv(self, path: Path) -> list[Payment]:
        """Parse a CSV statement file."""
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.DictReader(handle)
            return [
                self._build_payment(record, self._bank_name_for(path))
                for record in reader
                if self._has_content(record)
            ]

    def _parse_excel(self, path: Path) -> list[Payment]:
        """Parse an Excel workbook statement file."""
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            worksheet = workbook[workbook.sheetnames[0]]
            rows = list(worksheet.iter_rows(values_only=True))
        finally:
            workbook.close()

        if not rows:
            return []

        header = [self._stringify(cell).strip() for cell in rows[0]]
        result: list[Payment] = []
        for row in rows[1:]:
            if not row or all(cell is None or self._stringify(cell).strip() == "" for cell in row):
                continue
            record: dict[str, Any] = {}
            for index, value in enumerate(row):
                key = header[index] if index < len(header) else f"column_{index + 1}"
                record[key] = value
            result.append(self._build_payment(record, self._bank_name_for(path)))
        return result

    def _build_payment(self, record: Mapping[str, Any], source_bank: str) -> Payment:
        """Create a normalized Payment from a raw bank record."""
        normalized = self._normalize_record(record)
        return Payment(
            payment_date=normalized["payment_date"],
            amount=normalized["amount"],
            payer_name=normalized["payer_name"],
            description=normalized["description"],
            reference_number=normalized["reference_number"],
            source_bank=source_bank,
            original_record=dict(record),
        )

    def _normalize_record(self, record: Mapping[str, Any]) -> dict[str, Any]:
        """Normalize a raw record to the expected payment schema."""
        record_data = {self._stringify(key).lower(): value for key, value in record.items()}
        payment_date = self._normalize_date(
            self._first_present_value(
                record_data,
                (
                    "payment_date",
                    "date",
                    "transaction_date",
                    "posted_date",
                    "value_date",
                ),
            )
        )
        payer_name = self._normalize_name(
            self._first_present_value(
                record_data,
                ("payer_name", "payer", "name", "counterparty", "from", "beneficiary"),
            )
        )
        amount = self._normalize_amount(
            self._first_present_value(
                record_data,
                (
                    "amount",
                    "total",
                    "sum",
                    "value",
                    "debit",
                    "credit",
                    "payment_amount",
                ),
            )
        )
        description = self._normalize_description(
            self._first_present_value(
                record_data,
                (
                    "description",
                    "details",
                    "memo",
                    "narrative",
                    "message",
                    "info",
                ),
            )
        )
        reference_number = self._normalize_reference(
            self._first_present_value(
                record_data,
                (
                    "reference_number",
                    "reference",
                    "transaction_id",
                    "payment_reference",
                    "ref",
                ),
            )
        )

        return {
            "payment_date": payment_date,
            "amount": amount,
            "payer_name": payer_name,
            "description": description,
            "reference_number": reference_number,
        }

    @staticmethod
    def _bank_name_for(path: Path) -> str:
        """Produce a readable bank name from the file stem."""
        return path.stem.replace("_", " ").replace("-", " ").title()

    @staticmethod
    def _has_content(record: Mapping[str, Any]) -> bool:
        """Return whether a CSV row contains usable data."""
        return any(
            value is not None and str(value).strip() not in {"", "nan", "NaN"}
            for value in record.values()
        )

    @staticmethod
    def _first_present_value(record: Mapping[str, Any], keys: Sequence[str]) -> Any:
        """Return the first non-empty value found for the supported keys."""
        for key in keys:
            if key in record and record[key] is not None and str(record[key]).strip() != "":
                return record[key]
        return ""

    @staticmethod
    def _stringify(value: Any) -> str:
        """Convert a value to string safely."""
        if value is None:
            return ""
        if isinstance(value, str):
            return value
        return str(value)

    @classmethod
    def _normalize_date(cls, value: Any) -> date:
        """Normalize a payment date value to a date object."""
        if value is None:
            raise ValueError("Payment date is required.")

        text = cls._stringify(value).strip()
        if not text:
            raise ValueError("Payment date is required.")

        candidates = [
            "%Y-%m-%d",
            "%d.%m.%Y",
            "%d/%m/%Y",
            "%m/%d/%Y",
            "%Y/%m/%d",
            "%d-%m-%Y",
            "%Y.%m.%d",
            "%d %b %Y",
            "%d %B %Y",
        ]
        for pattern in candidates:
            try:
                return datetime.strptime(text, pattern).date()
            except ValueError:
                continue

        try:
            return datetime.fromisoformat(text.replace("Z", "+00:00")).date()
        except ValueError as exc:
            raise ValueError(f"Unsupported payment date: {value!r}") from exc

    @classmethod
    def _normalize_amount(cls, value: Any) -> Decimal:
        """Normalize the amount to a positive Decimal value."""
        if value is None:
            raise ValueError("Payment amount is required.")

        text = cls._stringify(value).strip()
        if not text:
            raise ValueError("Payment amount is required.")

        text = text.replace("€", "").replace("$", "").replace(" ", "")
        if text.startswith("(") and text.endswith(")"):
            text = f"-{text[1:-1]}"

        if "," in text and "." in text:
            if text.rfind(",") > text.rfind("."):
                text = text.replace(".", "").replace(",", ".")
            else:
                text = text.replace(",", "")
        elif "," in text:
            text = text.replace(",", ".")

        try:
            return abs(Decimal(text))
        except InvalidOperation as exc:
            raise ValueError(f"Unsupported payment amount: {value!r}") from exc

    @staticmethod
    def _normalize_name(value: Any) -> str:
        """Normalize payer names by collapsing whitespace and trimming edges."""
        text = " ".join(str(value).split()) if value is not None else ""
        return text

    @staticmethod
    def _normalize_description(value: Any) -> str:
        """Normalize description text by trimming and collapsing whitespace."""
        text = " ".join(str(value).split()) if value is not None else ""
        return text

    @staticmethod
    def _normalize_reference(value: Any) -> str | None:
        """Normalize a reference number to a string or None."""
        if value is None:
            return None
        text = " ".join(str(value).split())
        return text or None
