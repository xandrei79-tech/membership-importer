"""Match payments to customers using exact payer names and configured aliases."""

from collections.abc import Mapping
from pathlib import Path

from openpyxl import load_workbook

from ..models.customer import Customer
from ..models.match_result import MatchResult
from ..models.payment import Payment
from .customer_repository import CustomerRepository


class CustomerMatcher:
    """Match payment payer names to customers without allocating payments."""

    def __init__(
        self,
        repository: CustomerRepository,
        aliases: Mapping[str, str] | None = None,
        aliases_path: str | Path | None = None,
    ) -> None:
        self._repository = repository
        if aliases is not None:
            alias_records = aliases
        else:
            path = (
                Path(aliases_path)
                if aliases_path is not None
                else Path(__file__).resolve().parents[3] / "config" / "aliases.xlsx"
            )
            alias_records = self._load_aliases(path)
        self._aliases = {
            self._normalize(alias): self._normalize(canonical_name)
            for alias, canonical_name in alias_records.items()
            if self._normalize(alias) and self._normalize(canonical_name)
        }

    def match(self, payment: Payment) -> MatchResult:
        """Return the exact customer match, if one exists, for ``payment``."""
        payer_name = self._normalize(payment.payer_name)
        if not payer_name:
            return MatchResult(status="unmatched")

        lookup_name = self._aliases.get(payer_name, payer_name)
        candidates = tuple(
            customer
            for customer in self._repository.list_customers()
            if self._normalize(customer.customer_name) == lookup_name
        )

        if not candidates:
            return MatchResult(status="unmatched")
        if len(candidates) > 1:
            return MatchResult(status="ambiguous", candidates=candidates)
        return MatchResult(status="matched", customer=candidates[0], candidates=candidates)

    @staticmethod
    def _normalize(value: object) -> str:
        """Collapse whitespace and case-fold text for comparison."""
        return " ".join(str(value or "").split()).casefold()

    @staticmethod
    def _load_aliases(path: Path) -> dict[str, str]:
        """Load ``alias`` to ``canonical_name`` records from the config workbook."""
        if not path.exists():
            return {}

        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            rows = list(workbook.active.iter_rows(values_only=True))
        finally:
            workbook.close()

        if not rows:
            return {}

        headers = [str(value or "").strip().casefold() for value in rows[0]]
        try:
            alias_index = headers.index("alias")
            canonical_index = headers.index("canonical_name")
        except ValueError as exc:
            raise ValueError("Alias workbook must contain alias and canonical_name columns") from exc

        aliases: dict[str, str] = {}
        for row in rows[1:]:
            if row is None:
                continue
            alias = row[alias_index] if alias_index < len(row) else None
            canonical_name = row[canonical_index] if canonical_index < len(row) else None
            if alias is not None and canonical_name is not None:
                aliases[str(alias)] = str(canonical_name)
        return aliases